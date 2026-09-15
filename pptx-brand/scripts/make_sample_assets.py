#!/usr/bin/env python3
"""Generate the sample brand assets shipped with the pptx-brand skill.

Everything under ../assets/ is produced by this script from code, so the
sample brand carries no third-party images. Replace the outputs with the real
brand's files (same names, or update references/styles.json and assets.md)
and delete this script, or keep it to regenerate placeholders.

Usage:
    python make_sample_assets.py [--assets DIR] [--skip-template]

Needs Pillow and python-pptx (both present in the runtime image).
"""

from __future__ import annotations

import argparse
import importlib.util
import json
import shutil
import sys
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL_ROOT = HERE.parent
ENGINE_SCRIPTS = SKILL_ROOT.parent / "pptx" / "scripts"

NAVY = (11, 37, 69)
BLUE = (31, 111, 235)
ORANGE = (232, 89, 12)
TEAL = (15, 118, 110)
MINT = (15, 157, 138)
CHARCOAL = (31, 41, 51)
SLATE = (15, 23, 42)
WHITE = (255, 255, 255)
LIGHT = (243, 244, 246)
GOLD = (242, 201, 76)

ICON_NAMES = [
    "target", "chart-up", "shield", "gear",
    "users", "clock", "check", "warning",
    "cloud", "lock", "bolt", "flag",
    "document", "server", "globe", "bulb",
]


def _load_module(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    assert spec and spec.loader
    spec.loader.exec_module(module)
    return module


# --------------------------------------------------------------------------- #
# Drawing helpers
# --------------------------------------------------------------------------- #


def _font(size: int, bold: bool = True):
    from PIL import ImageFont

    candidates = [
        "C:/Windows/Fonts/segoeuib.ttf" if bold else "C:/Windows/Fonts/segoeui.ttf",
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).is_file():
            return ImageFont.truetype(candidate, size)
    try:
        return ImageFont.load_default(size=size)
    except TypeError:  # Pillow < 10.1
        return ImageFont.load_default()


def _mix(color_a, color_b, t: float):
    return tuple(int(round(a + (b - a) * t)) for a, b in zip(color_a, color_b))


def gradient(size, color_top, color_bottom, *, diagonal: bool = False):
    from PIL import Image, ImageDraw

    width, height = size
    image = Image.new("RGB", size, color_top)
    draw = ImageDraw.Draw(image)
    if diagonal:
        steps = width + height
        for i in range(steps):
            color = _mix(color_top, color_bottom, i / max(1, steps - 1))
            draw.line([(i, 0), (0, i)], fill=color)
        return image
    for y in range(height):
        draw.line([(0, y), (width, y)], fill=_mix(color_top, color_bottom, y / max(1, height - 1)))
    return image


def overlay_shapes(base, shapes):
    """Composite translucent shapes (drawn at 2x for smooth edges) onto an RGB image."""
    from PIL import Image, ImageDraw

    width, height = base.size
    layer = Image.new("RGBA", (width * 2, height * 2), (0, 0, 0, 0))
    draw = ImageDraw.Draw(layer)
    for kind, box, color, alpha in shapes:
        fill = (*color, int(255 * alpha))
        scaled = [v * 2 for v in box]
        if kind == "ellipse":
            draw.ellipse(scaled, fill=fill)
        elif kind == "polygon":
            draw.polygon([(scaled[i], scaled[i + 1]) for i in range(0, len(scaled), 2)], fill=fill)
        elif kind == "rect":
            draw.rectangle(scaled, fill=fill)
    layer = layer.resize((width, height), Image.LANCZOS)
    return Image.alpha_composite(base.convert("RGBA"), layer).convert("RGB")


def dot_grid(image, color, alpha: float, spacing: int = 48, radius: int = 2, region=None):
    from PIL import Image, ImageDraw

    width, height = image.size
    layer = Image.new("RGBA", image.size, (0, 0, 0, 0))
    draw = ImageDraw.Draw(layer)
    x0, y0, x1, y1 = region or (0, 0, width, height)
    for y in range(y0, y1, spacing):
        for x in range(x0, x1, spacing):
            draw.ellipse((x - radius, y - radius, x + radius, y + radius), fill=(*color, int(255 * alpha)))
    return Image.alpha_composite(image.convert("RGBA"), layer).convert("RGB")


# --------------------------------------------------------------------------- #
# Backgrounds (1920x1080 JPEG)
# --------------------------------------------------------------------------- #


def make_backgrounds(out_dir: Path) -> list[Path]:
    size = (1920, 1080)
    out_dir.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []

    # Formal cover: navy gradient, a pale diagonal band and a dot grid on the right.
    image = gradient(size, NAVY, (6, 20, 40), diagonal=True)
    image = overlay_shapes(image, [
        ("polygon", (1150, 0, 1920, 0, 1920, 1080, 780, 1080), WHITE, 0.05),
        ("polygon", (1450, 0, 1920, 0, 1920, 1080, 1180, 1080), BLUE, 0.18),
        ("ellipse", (1500, -300, 2300, 500), BLUE, 0.20),
    ])
    image = dot_grid(image, WHITE, 0.10, region=(1240, 620, 1920, 1080))
    written.append(_save_jpeg(image, out_dir / "cover-formal.jpg"))

    # Keynote cover: near-black with two glowing circles.
    image = gradient(size, SLATE, (3, 7, 18))
    image = overlay_shapes(image, [
        ("ellipse", (1200, 250, 2100, 1150), BLUE, 0.22),
        ("ellipse", (1450, 50, 1950, 550), (147, 197, 253), 0.10),
        ("ellipse", (-350, 600, 450, 1400), ORANGE, 0.16),
    ])
    written.append(_save_jpeg(image, out_dir / "cover-keynote.jpg"))

    # Section (light): white with pale geometric triangles on the right.
    image = gradient(size, WHITE, LIGHT)
    image = overlay_shapes(image, [
        ("polygon", (1300, 0, 1920, 0, 1920, 620), NAVY, 0.06),
        ("polygon", (1500, 1080, 1920, 1080, 1920, 520), BLUE, 0.10),
        ("polygon", (1650, 0, 1920, 0, 1920, 300), ORANGE, 0.55),
    ])
    written.append(_save_jpeg(image, out_dir / "section-light.jpg"))

    # Section (dark): deep navy with concentric arcs.
    image = gradient(size, (10, 33, 62), NAVY)
    image = overlay_shapes(image, [
        ("ellipse", (1250, 180, 2350, 1280), WHITE, 0.04),
        ("ellipse", (1400, 330, 2200, 1130), WHITE, 0.05),
        ("ellipse", (1550, 480, 2050, 980), ORANGE, 0.35),
    ])
    written.append(_save_jpeg(image, out_dir / "section-dark.jpg"))

    # Closing: navy to teal gradient with a radial highlight.
    image = gradient(size, NAVY, TEAL, diagonal=True)
    image = overlay_shapes(image, [
        ("ellipse", (600, 100, 1900, 1400), WHITE, 0.07),
        ("ellipse", (900, 300, 1600, 1000), WHITE, 0.06),
    ])
    written.append(_save_jpeg(image, out_dir / "closing.jpg"))

    # Training: light page with soft mint shapes; meant for dark text.
    image = gradient(size, WHITE, (236, 253, 245))
    image = overlay_shapes(image, [
        ("ellipse", (1350, -250, 2150, 550), MINT, 0.14),
        ("ellipse", (-250, 700, 550, 1500), TEAL, 0.10),
        ("rect", (0, 1040, 1920, 1080), TEAL, 0.9),
    ])
    written.append(_save_jpeg(image, out_dir / "training.jpg"))
    return written


def _save_jpeg(image, path: Path) -> Path:
    image.save(path, "JPEG", quality=82, optimize=True, progressive=True)
    return path


# --------------------------------------------------------------------------- #
# Icons: 16 line icons drawn with primitives on one sheet
# --------------------------------------------------------------------------- #


def _draw_icon(name: str, draw, color, stroke: int) -> None:
    c = 256
    line = {"fill": color, "width": stroke}
    outline = {"outline": color, "width": stroke}
    if name == "target":
        for r in (200, 130, 55):
            draw.ellipse((c - r, c - r, c + r, c + r), **outline)
    elif name == "chart-up":
        for i, h in enumerate((150, 240, 330)):
            x = 90 + i * 130
            draw.rectangle((x, 430 - h, x + 90, 430), fill=color)
        draw.line([(80, 470), (440, 470)], **line)
    elif name == "shield":
        draw.polygon([(256, 60), (440, 130), (430, 300), (256, 460), (82, 300), (72, 130)], **outline)
        draw.line([(180, 260), (240, 320), (340, 200)], **line)
    elif name == "gear":
        import math

        teeth = []
        for i in range(16):
            angle = math.pi * 2 * i / 16
            radius = 210 if i % 2 == 0 else 160
            teeth.append((c + radius * math.cos(angle), c + radius * math.sin(angle)))
        draw.polygon(teeth, fill=color)
        draw.ellipse((c - 80, c - 80, c + 80, c + 80), fill=(0, 0, 0, 0))
    elif name == "users":
        draw.ellipse((150, 90, 290, 230), **outline)
        draw.arc((80, 250, 360, 530), 180, 360, **line)
        draw.ellipse((320, 130, 420, 230), **outline)
        draw.arc((300, 250, 470, 450), 200, 340, **line)
    elif name == "clock":
        draw.ellipse((60, 60, 452, 452), **outline)
        draw.line([(256, 130), (256, 262), (350, 330)], **line)
    elif name == "check":
        draw.line([(90, 270), (210, 390), (430, 140)], **line)
    elif name == "warning":
        draw.polygon([(256, 70), (460, 430), (52, 430)], **outline)
        draw.line([(256, 190), (256, 320)], **line)
        draw.ellipse((238, 350, 274, 386), fill=color)
    elif name == "cloud":
        draw.ellipse((90, 220, 290, 420), fill=color)
        draw.ellipse((180, 130, 400, 350), fill=color)
        draw.ellipse((280, 220, 460, 400), fill=color)
        draw.rectangle((150, 320, 420, 420), fill=color)
    elif name == "lock":
        draw.arc((146, 60, 366, 280), 180, 360, **line)
        draw.line([(146, 170), (146, 240)], **line)
        draw.line([(366, 170), (366, 240)], **line)
        draw.rounded_rectangle((90, 240, 422, 460), radius=40, fill=color)
        draw.ellipse((232, 310, 280, 358), fill=(0, 0, 0, 0))
    elif name == "bolt":
        draw.polygon([(300, 40), (120, 300), (250, 300), (200, 480), (400, 210), (270, 210)], fill=color)
    elif name == "flag":
        draw.line([(110, 60), (110, 470)], **line)
        draw.polygon([(130, 80), (430, 80), (350, 190), (430, 300), (130, 300)], fill=color)
    elif name == "document":
        draw.polygon([(120, 60), (330, 60), (420, 150), (420, 460), (120, 460)], **outline)
        for y in (230, 300, 370):
            draw.line([(180, y), (360, y)], fill=color, width=stroke - 8)
    elif name == "server":
        for y in (70, 210, 350):
            draw.rounded_rectangle((70, y, 442, y + 110), radius=18, **outline)
            draw.ellipse((100, y + 40, 130, y + 70), fill=color)
    elif name == "globe":
        draw.ellipse((60, 60, 452, 452), **outline)
        draw.ellipse((170, 60, 342, 452), **outline)
        draw.line([(60, 256), (452, 256)], **line)
        draw.line([(256, 60), (256, 452)], **line)
    elif name == "bulb":
        draw.ellipse((110, 50, 402, 342), **outline)
        draw.rectangle((196, 320, 316, 400), fill=color)
        draw.rectangle((216, 410, 296, 450), fill=color)


def make_icon_sheet(path: Path, color, *, cell: int = 256, gap: int = 32, cols: int = 4) -> Path:
    from PIL import Image, ImageDraw

    rows = (len(ICON_NAMES) + cols - 1) // cols
    sheet = Image.new("RGBA", (cols * cell + (cols + 1) * gap, rows * cell + (rows + 1) * gap), (0, 0, 0, 0))
    for index, name in enumerate(ICON_NAMES):
        canvas = Image.new("RGBA", (512, 512), (0, 0, 0, 0))
        _draw_icon(name, ImageDraw.Draw(canvas), (*color, 255), 36)
        icon = canvas.resize((cell - 24, cell - 24), Image.LANCZOS)
        x = gap + (index % cols) * (cell + gap) + 12
        y = gap + (index // cols) * (cell + gap) + 12
        sheet.paste(icon, (x, y), icon)
    path.parent.mkdir(parents=True, exist_ok=True)
    sheet.save(path, optimize=True)
    return path


def recolor_icons(src_dir: Path, dst_dir: Path, color) -> int:
    from PIL import Image

    dst_dir.mkdir(parents=True, exist_ok=True)
    count = 0
    for icon_path in sorted(src_dir.glob("*.png")):
        if icon_path.name == "contact-sheet.png":
            continue
        icon = Image.open(icon_path).convert("RGBA")
        alpha = icon.getchannel("A")
        tinted = Image.new("RGBA", icon.size, (*color, 255))
        tinted.putalpha(alpha)
        tinted.save(dst_dir / icon_path.name, optimize=True)
        count += 1
    return count


# --------------------------------------------------------------------------- #
# Logo
# --------------------------------------------------------------------------- #


def make_logo(path: Path, color, *, text_color=None) -> Path:
    from PIL import Image, ImageDraw

    text_color = text_color or color
    canvas = Image.new("RGBA", (2400, 720), (0, 0, 0, 0))
    draw = ImageDraw.Draw(canvas)
    draw.rounded_rectangle((60, 60, 660, 660), radius=120, fill=(*color, 255))
    draw.polygon([(210, 200), (520, 200), (520, 280), (300, 280), (300, 320), (480, 320), (480, 400), (300, 400), (300, 440), (520, 440), (520, 520), (210, 520)], fill=(255, 255, 255, 255) if color != WHITE else (*NAVY, 255))
    draw.text((760, 150), "EFP", font=_font(360), fill=(*text_color, 255))
    draw.text((770, 540), "ENGINEERING FLOW PLATFORM", font=_font(70, bold=False), fill=(*text_color, 220))
    logo = canvas.resize((1200, 360), Image.LANCZOS)
    path.parent.mkdir(parents=True, exist_ok=True)
    logo.save(path, optimize=True)
    return path


# --------------------------------------------------------------------------- #
# Template deck: text examples in the formal style, theme XML patched
# --------------------------------------------------------------------------- #

THEME_COLORS = {
    "dk1": "1F2933", "lt1": "FFFFFF", "dk2": "0B2545", "lt2": "F3F4F6",
    "accent1": "1F6FEB", "accent2": "E8590C", "accent3": "0F9D8A", "accent4": "8E44AD",
    "accent5": "C0392B", "accent6": "F2C94C", "hlink": "1F6FEB", "folHlink": "6B7280",
}
THEME_FONTS = {"latin": "Calibri", "ea": "Microsoft YaHei"}


def template_spec(assets: Path) -> dict:
    icons = assets / "icons" / "navy"
    return {
        "title": "EFP 企业演示模板",
        "subtitle": "文字示例与版式参考 · 正式内容请用 /pptx-brand 生成",
        "author": "平台团队",
        "date": "2026-09",
        "footer": "EFP · 内部资料",
        "theme": {"primary": "0B2545", "accent": "E8590C", "east_asian_font": "Microsoft YaHei"},
        "slides": [
            {"type": "bullets", "title": "议程页：一句话说明本次汇报要回答的问题", "bullets": [
                "现状：采用率与稳定性", "本季度交付了什么", "成本与用量趋势", "风险与下季度诉求"],
                "notes": "议程页控制在一分钟内，标题写成本次汇报要回答的问题。"},
            {"type": "section", "title": "分节页：用一个短语概括本节", "subtitle": "副标题补一句本节的结论"},
            {"type": "bullets", "title": "要点页：标题是结论，不是主题词", "bullets": [
                "每页最多六条，每条一句话，二十字左右",
                "先说结论，再说依据；数字进 KPI 页或图表页",
                {"text": "二级要点只用来补充证据或例子", "level": 1},
                "动词开头，避免形容词堆砌",
                "细节写进备注，不要塞进页面"],
                "notes": "备注里放讲稿或数据来源，页面只留结论。"},
            {"type": "two_column", "title": "对比页：左右各说一件事", "left": {"heading": "现状", "bullets": [
                "手工整理，平均两天", "口径不一致"]}, "right": {"heading": "目标", "bullets": [
                "助手自动生成，十分钟", "同一份口径，可追溯"]}},
            {"type": "kpi", "title": "数字页：四个以内的关键指标", "metrics": [
                {"value": "412", "label": "周活成员", "description": "较上季 +38%", "icon": str(icons / "users.png")},
                {"value": "97.8%", "label": "对话成功率", "icon": str(icons / "check.png")},
                {"value": "3.1 秒", "label": "首字延迟中位数", "icon": str(icons / "clock.png")},
                {"value": "26", "label": "生产助手数", "icon": str(icons / "server.png")}]},
            {"type": "cards", "title": "卡片页：三到四个并列的要点", "items": [
                {"icon": str(icons / "target.png"), "title": "目标明确", "text": "每页只回答一个问题，标题即结论。"},
                {"icon": str(icons / "shield.png"), "title": "口径一致", "text": "数字来自同一数据源，注明时间范围。"},
                {"icon": str(icons / "bolt.png"), "title": "快速迭代", "text": "先出第一版再补细节，改 spec 重建即可。"}]},
            {"type": "chart", "title": "图表页：标题说明趋势，图表只放一组数据", "chart": {
                "kind": "column", "title": "周活成员", "categories": ["7 月", "8 月", "9 月"],
                "series": [{"name": "周活", "values": [321, 360, 412]}], "data_labels": True}},
            {"type": "table", "title": "表格页：不超过六行五列", "headers": ["供应商", "请求数", "成本"], "rows": [
                ["GitHub Copilot", "184,200", "$9,120"], ["AI Platform", "62,400", "$4,368"], ["合计", "246,600", "$13,488"]]},
            {"type": "quote", "quote": "引用页：一句真实的客户或干系人原话，注明出处。", "attribution": "支付小组技术负责人"},
            {"type": "closing", "title": "谢谢", "subtitle": "问题请到 #efp-platform"},
        ],
    }


def patch_theme(pptx_path: Path) -> None:
    """Rewrite ppt/theme/theme*.xml inside the package with the brand colours and fonts."""
    from lxml import etree

    ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    tmp_path = pptx_path.with_suffix(".tmp")
    with zipfile.ZipFile(pptx_path) as src, zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename.startswith("ppt/theme/") and item.filename.endswith(".xml"):
                root = etree.fromstring(data)
                scheme = root.find(f".//{{{ns}}}clrScheme")
                if scheme is not None:
                    scheme.set("name", "EFP")
                    for slot, value in THEME_COLORS.items():
                        node = scheme.find(f"{{{ns}}}{slot}")
                        if node is None:
                            continue
                        for child in list(node):
                            node.remove(child)
                        node.append(etree.SubElement(node, f"{{{ns}}}srgbClr", val=value))
                for kind in ("majorFont", "minorFont"):
                    font = root.find(f".//{{{ns}}}fontScheme/{{{ns}}}{kind}")
                    if font is None:
                        continue
                    for script, face in THEME_FONTS.items():
                        node = font.find(f"{{{ns}}}{script}")
                        if node is not None:
                            node.set("typeface", face)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            dst.writestr(item, data)
    shutil.move(tmp_path, pptx_path)


def add_master_footer_bar(pptx_path: Path) -> bool:
    """Put a thin accent bar on the slide master so every slide drawn on it inherits the decoration."""
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.shapes.autoshape import Shape
        from pptx.util import Inches

        prs = Presentation(str(pptx_path))
        master = prs.slide_master
        tree = master.shapes._spTree
        shape_id = master.shapes._next_shape_id
        sp = tree.add_autoshape(shape_id, "Brand footer bar", "rect", 0, prs.slide_height - Inches(0.12), prs.slide_width, Inches(0.12))
        shape = Shape(sp, master.shapes)
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string("0B2545")
        shape.line.fill.background()
        sp2 = tree.add_autoshape(shape_id + 1, "Brand footer accent", "rect", 0, prs.slide_height - Inches(0.12), Inches(1.6), Inches(0.12))
        accent = Shape(sp2, master.shapes)
        accent.fill.solid()
        accent.fill.fore_color.rgb = RGBColor.from_string("E8590C")
        accent.line.fill.background()
        prs.save(str(pptx_path))
        return True
    except Exception as exc:  # the template is still usable without the bar
        print(f"note: master footer bar skipped ({exc.__class__.__name__}: {exc})", file=sys.stderr)
        return False


def make_template(assets: Path, build_deck) -> Path:
    out = assets / "template" / "efp-brand-template.pptx"
    out.parent.mkdir(parents=True, exist_ok=True)
    spec = template_spec(assets)
    summary = build_deck.build_deck(spec, out, base_dir=assets)
    if not summary.get("ok"):
        raise SystemExit(f"template build failed: {summary}")
    patch_theme(out)
    add_master_footer_bar(out)
    return out


# --------------------------------------------------------------------------- #


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--assets", default=str(SKILL_ROOT / "assets"), help="Output directory (default: ../assets)")
    parser.add_argument("--skip-template", action="store_true", help="Do not rebuild the template deck")
    args = parser.parse_args(argv)
    assets = Path(args.assets).resolve()

    slice_icons = _load_module("slice_icons", ENGINE_SCRIPTS / "slice_icons.py")
    build_deck = _load_module("build_deck", ENGINE_SCRIPTS / "build_deck.py")

    report: dict = {"assets": str(assets)}
    report["backgrounds"] = [p.name for p in make_backgrounds(assets / "backgrounds")]
    sheet = make_icon_sheet(assets / "icons" / "icon-sheet.png", NAVY)
    sliced = slice_icons.slice_sheet(sheet, assets / "icons" / "navy", names=ICON_NAMES, size=256)
    if not sliced.get("ok"):
        raise SystemExit(f"icon slicing failed: {sliced}")
    report["icons"] = sliced["names"]
    report["white_icons"] = recolor_icons(assets / "icons" / "navy", assets / "icons" / "white", WHITE)
    report["logos"] = [
        make_logo(assets / "logo" / "logo-primary.png", NAVY).name,
        make_logo(assets / "logo" / "logo-white.png", WHITE, text_color=WHITE).name,
    ]
    if not args.skip_template:
        report["template"] = make_template(assets, build_deck).name
    total = sum(p.stat().st_size for p in assets.rglob("*") if p.is_file())
    report["total_bytes"] = total
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
