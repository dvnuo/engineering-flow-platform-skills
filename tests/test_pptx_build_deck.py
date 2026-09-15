"""Behaviour of pptx/scripts/build_deck.py.

Spec validation and the CLI contract need only the standard library; the
rendering test runs when python-pptx happens to be installed.
"""

from __future__ import annotations

import importlib.util
import json
from pathlib import Path

import pytest

from tests._skill_presence import require_skill


REPO_ROOT = Path(__file__).resolve().parents[1]
SCRIPT = REPO_ROOT / "pptx" / "scripts" / "build_deck.py"
EXAMPLE = REPO_ROOT / "pptx" / "references" / "spec-example.json"


def _python_pptx_installed() -> bool:
    """True only for the real library.

    With the repo root on sys.path the skill directory ``pptx/`` is importable
    as an empty namespace package, so ``find_spec("pptx")`` is not enough.
    """
    try:
        module = importlib.import_module("pptx")
    except ImportError:
        return False
    return hasattr(module, "Presentation")


needs_python_pptx = pytest.mark.skipif(not _python_pptx_installed(), reason="python-pptx is not installed")


def _load_module():
    require_skill("pptx")
    spec = importlib.util.spec_from_file_location("pptx_build_deck", SCRIPT)
    module = importlib.util.module_from_spec(spec)
    assert spec and spec.loader
    spec.loader.exec_module(module)
    return module


def _example() -> dict:
    return json.loads(EXAMPLE.read_text(encoding="utf-8"))


def test_example_spec_validates_clean() -> None:
    module = _load_module()

    errors, warnings = module.validate_spec(_example(), EXAMPLE.parent)

    assert errors == []
    assert warnings == []


def test_validate_spec_reports_errors_and_warnings(tmp_path: Path) -> None:
    module = _load_module()
    spec = {
        "title": "Broken deck",
        "theme": {"primary": "not-a-colour"},
        "slides": [
            {"type": "chart", "title": "t", "chart": {"kind": "radar", "categories": ["a"], "series": [{"values": [1, 2]}]}},
            {"type": "image", "title": "i", "image": "missing.png"},
            {"type": "bullets", "title": "b", "bullets": [{"text": "x", "level": 5}] + [f"bullet {i}" for i in range(9)]},
            {"type": "kpi", "title": "k", "metrics": [{"value": "1"}]},
            {"type": "quote"},
            {"type": "unknown", "title": "u"},
            "not-an-object",
        ],
    }

    errors, warnings = module.validate_spec(spec, tmp_path)

    joined = "\n".join(errors)
    assert "chart.kind must be one of" in joined
    assert "has 2 values for 1 categories" in joined
    assert "image not found: missing.png" in joined
    assert "level must be 0, 1 or 2" in joined
    assert "metrics[0] needs 'value' and 'label'" in joined
    assert "'quote' text is required" in joined
    assert "unknown type 'unknown'" in joined
    assert "slides[6]: must be an object" in joined

    joined_warnings = "\n".join(warnings)
    assert "theme.primary is not a 6-digit hex colour" in joined_warnings
    assert "10 bullets on one slide" in joined_warnings


def test_validate_spec_rejects_non_object_and_empty_slides() -> None:
    module = _load_module()

    assert module.validate_spec([], None)[0] == ["spec must be a JSON object"]
    errors, _ = module.validate_spec({"title": "x", "slides": []}, None)
    assert "'slides' must be a non-empty list" in errors


def test_resolve_output_path_defaults_and_extension() -> None:
    module = _load_module()

    assert module.resolve_output_path({"title": "Q3 Platform Review"}, None) == Path("output/Q3-Platform-Review.pptx")
    assert module.resolve_output_path({"title": "季度复盘 / 汇报"}, None) == Path("output/季度复盘-汇报.pptx")
    assert module.resolve_output_path({}, None) == Path("output/deck.pptx")
    assert module.resolve_output_path({"output": "output/x.ppt"}, None) == Path("output/x.pptx")
    assert module.resolve_output_path({"output": "output/x.pptx"}, "custom/y.pptx") == Path("custom/y.pptx")


def test_resolve_theme_keeps_valid_overrides_only() -> None:
    module = _load_module()

    theme = module.resolve_theme({"theme": {"primary": "#123abc", "accent": "nope", "east_asian_font": "Noto Sans CJK SC"}})

    assert theme["primary"] == "123ABC"
    assert theme["accent"] == module.DEFAULT_THEME["accent"]
    assert theme["east_asian_font"] == "Noto Sans CJK SC"
    assert theme["title_font"] == theme["font"]


def test_cli_validate_only_prints_json(capsys: pytest.CaptureFixture[str]) -> None:
    module = _load_module()

    exit_code = module.main([str(EXAMPLE), "--validate-only", "--quiet"])

    payload = json.loads(capsys.readouterr().out.strip())
    assert exit_code == 0
    assert payload["ok"] is True
    assert payload["validated"] is True
    assert payload["slides"] == len(_example()["slides"])


def test_cli_reports_invalid_spec_with_exit_code_2(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    module = _load_module()
    spec_path = tmp_path / "bad.json"
    spec_path.write_text(json.dumps({"title": "x", "slides": [{"type": "quote"}]}), encoding="utf-8")

    exit_code = module.main([str(spec_path), "--quiet"])

    payload = json.loads(capsys.readouterr().out.strip())
    assert exit_code == 2
    assert payload["ok"] is False
    assert any("'quote' text is required" in error for error in payload["errors"])


def test_cli_print_example_matches_reference(capsys: pytest.CaptureFixture[str]) -> None:
    module = _load_module()

    assert module.main(["--print-example"]) == 0
    assert json.loads(capsys.readouterr().out) == _example()


@needs_python_pptx
def test_build_example_deck_renders_every_slide(tmp_path: Path) -> None:
    module = _load_module()
    output = tmp_path / "out" / "deck.pptx"

    summary = module.build_deck(_example(), output, base_dir=EXAMPLE.parent)

    assert summary["ok"] is True
    assert summary["warnings"] == []
    # The example has no explicit title slide, so one is generated first.
    assert summary["slides"] == len(_example()["slides"]) + 1
    assert output.is_file() and output.stat().st_size > 0

    from pptx import Presentation

    presentation = Presentation(str(output))
    slides = list(presentation.slides)
    assert len(slides) == summary["slides"]
    first_texts = [shape.text_frame.text for shape in slides[0].shapes if shape.has_text_frame]
    assert "Q3 Platform Review" in first_texts
    agenda = slides[1]
    assert agenda.has_notes_slide
    assert "under a minute" in agenda.notes_slide.notes_text_frame.text
    assert presentation.core_properties.title == "Q3 Platform Review"


@needs_python_pptx
def test_build_keeps_every_shape_inside_the_slide(tmp_path: Path) -> None:
    """Layout maths is relative to the slide size; nothing may hang off the edge."""
    module = _load_module()
    spec = _example()
    spec["theme"]["east_asian_font"] = "Microsoft YaHei"
    spec["slides"].insert(
        1,
        {
            "type": "kpi",
            "title": "六个指标",
            "metrics": [{"value": str(i), "label": f"指标 {i}"} for i in range(6)],
        },
    )
    output = tmp_path / "bounds.pptx"

    summary = module.build_deck(spec, output, base_dir=EXAMPLE.parent)
    assert summary["ok"] is True

    from pptx import Presentation

    presentation = Presentation(str(output))
    width, height = presentation.slide_width, presentation.slide_height
    for index, slide in enumerate(presentation.slides, start=1):
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0, (index, shape.name)
            assert shape.left + shape.width <= width, (index, shape.name)
            assert shape.top + shape.height <= height, (index, shape.name)


# --------------------------------------------------------------------------- #
# Styles, backgrounds, logos, icons and cards
# --------------------------------------------------------------------------- #

BRAND_DIR = REPO_ROOT / "pptx-brand"
BRAND_STYLES = BRAND_DIR / "references" / "styles.json"


def _brand_present() -> bool:
    return BRAND_STYLES.is_file() and (BRAND_DIR / "assets" / "template" / "efp-brand-template.pptx").is_file()


def test_validate_cards_kpi_icons_and_backgrounds(tmp_path: Path) -> None:
    module = _load_module()
    icon = tmp_path / "icon.png"
    icon.write_bytes(b"png")
    picture = tmp_path / "bg.jpg"
    picture.write_bytes(b"jpg")
    spec = {
        "title": "Assets",
        "logo": {"light": "icon.png", "position": "top-right", "hero_position": "bottom-left", "width_in": 1.2},
        "defaults": {
            "backgrounds": {"title": "bg.jpg", "section": {"image": "bg.jpg", "text": "dark", "overlay": 0}},
            "overlay": 0.4,
        },
        "slides": [
            {"type": "cards", "title": "c", "items": [{"title": "One", "text": "t", "icon": "icon.png"}, {"title": "Two"}]},
            {"type": "kpi", "title": "k", "metrics": [{"value": "1", "label": "a", "icon": "icon.png"}]},
            {"type": "section", "title": "s", "background": {"image": "bg.jpg", "overlay": 0.3, "overlay_color": "112233"}},
            {"type": "quote", "quote": "q", "background": ""},
            {"type": "closing", "background": "bg.jpg"},
        ],
    }

    errors, warnings = module.validate_spec(spec, tmp_path)

    assert errors == []
    assert warnings == []


def test_validate_reports_bad_assets_and_placements(tmp_path: Path) -> None:
    module = _load_module()
    spec = {
        "title": "Broken assets",
        "logo": {"position": "middle"},
        "template": "missing.potx",
        "defaults": {"backgrounds": {"bullets": "x.jpg"}, "overlay": 2},
        "slides": [
            {"type": "cards", "title": "c", "items": [{"text": "no title"}, {"title": "x", "icon": "nope.png"}]},
            {"type": "cards", "title": "too many", "items": [{"title": str(i)} for i in range(9)]},
            {"type": "kpi", "title": "k", "metrics": [{"value": "1", "label": "a", "icon": "nope.png"}]},
            {"type": "bullets", "title": "b", "bullets": ["x"], "background": "bg.jpg"},
            {"type": "section", "title": "s", "background": {"image": "missing.jpg", "overlay": 1.5, "text": "grey"}},
        ],
    }

    errors, _ = module.validate_spec(spec, tmp_path)

    joined = "\n".join(errors)
    assert "logo needs at least one of 'light', 'dark' or 'image'" in joined
    assert "logo.position must be one of" in joined
    assert "template not found: missing.potx" in joined
    assert "defaults.backgrounds.bullets" in joined
    assert "defaults.overlay must be a number between 0 and 1" in joined
    assert "items[0]: needs a 'title'" in joined
    assert "icon not found: nope.png" in joined
    assert "at most 8 cards per slide" in joined
    assert "'background' is only supported on" in joined
    assert "background.overlay must be a number between 0 and 1" in joined
    assert "background.text must be 'light' or 'dark'" in joined
    assert "background image not found: missing.jpg" in joined


def test_apply_style_lets_spec_values_win() -> None:
    module = _load_module()
    style = {
        "theme": {"primary": "111111", "accent": "222222"},
        "template": "t.pptx",
        "footer": "style footer",
        "logo": {"light": "l.png"},
        "backgrounds": {"title": "a.jpg", "closing": "c.jpg"},
        "overlay": 0.4,
    }
    spec = {
        "title": "x",
        "theme": {"accent": "ABCDEF", "text": ""},
        "footer": "spec footer",
        "defaults": {"backgrounds": {"title": "mine.jpg"}, "overlay_color": "000000"},
        "slides": [],
    }

    merged = module.apply_style(spec, style)

    assert merged["theme"] == {"primary": "111111", "accent": "ABCDEF"}
    assert merged["footer"] == "spec footer"
    assert merged["template"] == "t.pptx"
    assert merged["logo"] == {"light": "l.png"}
    assert merged["defaults"] == {
        "backgrounds": {"title": "mine.jpg", "closing": "c.jpg"},
        "overlay": 0.4,
        "overlay_color": "000000",
    }
    # The input spec is not mutated.
    assert spec["defaults"] == {"backgrounds": {"title": "mine.jpg"}, "overlay_color": "000000"}


def test_load_and_select_styles(tmp_path: Path) -> None:
    module = _load_module()
    styles_path = tmp_path / "styles.json"
    styles_path.write_text(
        json.dumps({"default_style": "a", "styles": {"a": {"theme": {"primary": "111111"}}, "b": {"footer": "B", "bogus": 1}}}),
        encoding="utf-8",
    )

    document, errors = module.load_styles(styles_path)
    assert document["default_style"] == "a"
    assert errors == [f"style 'b' has unknown keys ['bogus']; allowed: {sorted(module.STYLE_KEYS)}"]

    name, style, errors = module.select_style(document, None)
    assert (name, style, errors) == ("a", {"theme": {"primary": "111111"}}, [])
    name, style, errors = module.select_style(document, "zzz")
    assert style is None and "unknown style 'zzz'" in errors[0]
    assert module.load_styles(tmp_path / "missing.json")[1][0].startswith("styles file unreadable")
    (tmp_path / "empty.json").write_text("{}", encoding="utf-8")
    assert "non-empty 'styles' object" in module.load_styles(tmp_path / "empty.json")[1][0]


def test_resolve_theme_derives_hero_colours() -> None:
    module = _load_module()

    theme = module.resolve_theme({"theme": {"primary": "0B2545", "background": "0F172A", "light": "1E293B"}})

    assert theme["hero"] == "0B2545"
    assert theme["hero_text"] == "0F172A"
    assert theme["hero_subtext"] == "1E293B"
    assert module.resolve_theme({"theme": {"hero_text": "ffffff"}})["hero_text"] == "FFFFFF"


def test_cli_list_styles_and_style_errors(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    module = _load_module()
    styles_path = tmp_path / "styles.json"
    styles_path.write_text(
        json.dumps({"brand": "Acme", "default_style": "x", "styles": {"x": {"description": "d"}}}), encoding="utf-8"
    )

    assert module.main(["--list-styles", "--styles", str(styles_path), "--quiet"]) == 0
    payload = json.loads(capsys.readouterr().out.strip())
    assert payload == {"ok": True, "brand": "Acme", "default_style": "x", "styles": {"x": "d"}}

    spec_path = tmp_path / "spec.json"
    spec_path.write_text(json.dumps({"title": "t", "style": "nope", "slides": [{"type": "closing"}]}), encoding="utf-8")
    assert module.main([str(spec_path), "--styles", str(styles_path), "--validate-only", "--quiet"]) == 2
    assert "unknown style 'nope'" in capsys.readouterr().out
    assert module.main([str(spec_path), "--validate-only", "--quiet"]) == 2
    assert "no styles file was given" in capsys.readouterr().out


@needs_python_pptx
@pytest.mark.skipif(not _brand_present(), reason="pptx-brand skill is not on this branch")
@pytest.mark.parametrize("style", ["formal", "review", "keynote", "training"])
def test_build_brand_styles_with_backgrounds_logo_icons_and_cards(tmp_path: Path, style: str) -> None:
    module = _load_module()
    icons = BRAND_DIR / "assets" / "icons" / ("white" if style == "keynote" else "navy")
    spec = {
        "title": "品牌风格样例",
        "subtitle": "四种风格同一份内容",
        "style": style,
        "styles_file": str(BRAND_STYLES),
        "slides": [
            {"type": "kpi", "title": "数字页", "metrics": [
                {"value": "412", "label": "周活", "icon": str(icons / "users.png")},
                {"value": "97.8%", "label": "成功率", "icon": str(icons / "check.png")},
            ]},
            {"type": "section", "title": "分节页", "subtitle": "副标题"},
            {"type": "cards", "title": "卡片页", "items": [
                {"icon": str(icons / "target.png"), "title": "目标", "text": "一句说明"},
                {"icon": str(icons / "shield.png"), "title": "可靠"},
                {"title": "无图标卡片", "text": "也可以"},
            ]},
            {"type": "quote", "quote": "一句原话", "attribution": "某人"},
            {"type": "closing", "title": "谢谢"},
        ],
    }
    spec_path = tmp_path / "brand.json"
    spec_path.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
    output = tmp_path / f"{style}.pptx"

    assert module.main([str(spec_path), "--output", str(output), "--quiet"]) == 0

    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(str(output))
    slides = list(presentation.slides)
    assert len(slides) == 6
    width, height = presentation.slide_width, presentation.slide_height
    pictures = 0
    for index, slide in enumerate(slides, start=1):
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0, (style, index, shape.name)
            assert shape.left + shape.width <= width and shape.top + shape.height <= height, (style, index, shape.name)
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pictures += 1
    # Logo on every slide, two KPI icons, two card icons, and a background on the hero slides.
    assert pictures >= 6 + 4 + 2, (style, pictures)
    cover = [shape for shape in slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.width == width]
    assert cover and cover[0].height == height
